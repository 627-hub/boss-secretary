"""Excel 责任矩阵导入器（PRD §7.4-7.6）。

Sheet1_责任矩阵 / Sheet2_枚举字典 / Sheet3_变更记录 三表结构；
读入 → 编译为决策表 YAML → 语法/枚举/重叠校验 → 干跑回放 → commit 落库
（matrix_versions 激活+旧版退役 + unified diff 写 audit）。

CLI（经 boss_secretary.matrix）:
  ... matrix import config/matrix/reimburse_v1.xlsx --yes
  ... matrix export config/matrix/reimburse_v1.yaml --xlsx config/matrix/reimburse_v1.xlsx
  ... matrix commit config/matrix/reimburse_v1.xlsx --db data/secretary.db --operator boss
"""
from __future__ import annotations

import difflib
import hashlib
import itertools
import re
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook

from boss_secretary.core import matrix as M
from boss_secretary import models as MD

SHEET_RULES = "责任矩阵"
SHEET_ENUM = "枚举字典"
SHEET_LOG = "变更记录"

SEC_INPUT = "[输入列定义]"
SEC_OUTPUT = "[输出列]"
SEC_MAP = "[动作→审批人映射]"

PRIORITY_COL = "优先级"
ACTION_COL = "审批动作"
NOTIFY_COL = "抄送"

ID_COL = "规则ID"

_TITLE_RE = re.compile(r"matrix:\s*(\S+)\s*\|.*hit_policy:\s*(\S+)", re.S)

DEFAULT_ACTIONS = ["AUTO_APPROVE", "MANAGER", "MANAGER_BOSS",
                   "MANAGER_BOSS_FINANCE_CONSIGN", "MANUAL_REVIEW", "AUTO_REJECT"]
DEFAULT_NOTIFY = ["finance", "employee", "audit"]


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _split_list(s: str) -> list[str]:
    return [x.strip() for x in re.split(r"[,，、]", s or "") if x.strip()]


def _parse_enum_sheet(ws):
    inputs: list[dict] = []
    actions: list[str] = []
    notify_targets: list[str] = []
    mapping: dict[str, str] = {}
    section = None
    for row in ws.iter_rows(values_only=True):
        vals = [_cell_str(v) for v in row] if row else []
        if not vals or all(v == "" for v in vals):
            continue
        c0 = vals[0]
        if c0.startswith("["):
            section = c0
            continue
        if section == SEC_INPUT:
            if c0 == "列名":
                continue
            if not (len(vals) > 1 and vals[1]):
                raise M.MatrixLintError(f"输入列 {c0} 缺少类型（string/number）")
            inputs.append({"name": c0,
                           "type": vals[1].lower(),
                           "enum": _split_list(vals[2]) if len(vals) > 2 else [],
                           "header": vals[3] if len(vals) > 3 else ""})
        elif section == SEC_OUTPUT:
            if c0.startswith("动作"):
                actions = _split_list(vals[1] if len(vals) > 1 else "")
            elif c0.startswith("抄送"):
                notify_targets = _split_list(vals[1] if len(vals) > 1 else "")
        elif section == SEC_MAP:
            if c0 and len(vals) > 1:
                mapping[c0] = vals[1]
    return inputs, actions, notify_targets, mapping


def read_xlsx(path: str | Path) -> dict:
    p = Path(path)
    if not p.exists():
        raise M.MatrixLintError(f"文件不存在: {p}")
    wb = load_workbook(p, data_only=True, read_only=True)
    missing = [s for s in (SHEET_RULES, SHEET_ENUM, SHEET_LOG) if s not in wb.sheetnames]
    if missing:
        raise M.MatrixLintError(f"缺少工作表: {missing}（可用 export 生成模板）")

    ws = wb[SHEET_RULES]
    title = _cell_str(ws.cell(row=1, column=1).value)
    tmatch = _TITLE_RE.search(title)
    if not tmatch:
        raise M.MatrixLintError(f"{SHEET_RULES}!A1 需为 'matrix: 名称 | hit_policy: FIRST' 格式")
    name, policy = tmatch.group(1), tmatch.group(2).upper()
    if policy not in (M.FIRST, M.UNIQUE):
        raise M.MatrixLintError(f"hit_policy 非法: {policy}")

    header_row, headers = None, []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [_cell_str(v) for v in row]
        if vals and vals[0] == PRIORITY_COL:
            header_row, headers = i, vals
            break
    if header_row is None:
        raise M.MatrixLintError(f"{SHEET_RULES} 缺少表头行（首列={PRIORITY_COL}）")

    inputs, actions, notify_targets, mapping = _parse_enum_sheet(wb[SHEET_ENUM])
    if not inputs:
        raise M.MatrixLintError(f"{SHEET_ENUM} 缺少 {SEC_INPUT} 输入列定义")

    rules_raw: list[dict] = []
    rule_rows: list[int] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= header_row:
            continue
        vals = [_cell_str(v) for v in row]
        if all(v == "" for v in vals):
            continue
        rec = dict(zip(headers, vals))
        if not rec.get(ACTION_COL):
            raise M.MatrixLintError(f"{SHEET_RULES} 第 {i} 行缺少 {ACTION_COL}")
        rules_raw.append(rec)
        rule_rows.append(i)

    version, changed_by, summary = None, "", ""
    for row in wb[SHEET_LOG].iter_rows(min_row=2, values_only=True):
        vals = [_cell_str(v) for v in row] if row else []
        if not vals or not vals[0] or vals[0] == "version":
            continue
        try:
            version = int(float(vals[0]))
        except ValueError:
            continue
        changed_by = vals[2] if len(vals) > 2 else ""
        summary = vals[3] if len(vals) > 3 else ""
    if version is None:
        raise M.MatrixLintError(f"{SHEET_LOG} 缺少 version 行")

    return {"name": name, "hit_policy": policy, "headers": headers, "inputs": inputs,
            "actions": actions, "notify_targets": notify_targets, "mapping": mapping,
            "rules_raw": rules_raw, "rule_rows": rule_rows, "version": version,
            "changed_by": changed_by, "summary": summary}


def xlsx_to_matrix_dict(px: dict) -> dict:
    header_of: dict[str, str] = {}
    inputs_compiled: list[dict] = []
    for d in px["inputs"]:
        h = d.get("header") or d["name"]
        header_of[h] = d["name"]
        item: dict[str, Any] = {"name": d["name"], "type": d["type"]}
        if d["enum"]:
            item["enum"] = d["enum"]
        if d.get("header"):
            item["header"] = d["header"]
        inputs_compiled.append(item)

    rules = []
    for idx, rec in enumerate(px["rules_raw"], 1):
        cells: dict[str, Any] = {}
        for h, v in rec.items():
            if h in (PRIORITY_COL, ID_COL, ACTION_COL, NOTIFY_COL):
                continue
            col = header_of.get(h)
            if col is None:
                raise M.MatrixLintError(f"表头「{h}」未在 {SEC_INPUT} 中定义")
            cells[col] = v if v != "" else None
        prio = rec.get(PRIORITY_COL, "")
        rid_raw = rec.get(ID_COL, "")
        if rid_raw:
            rid = rid_raw
        else:
            try:
                rid = f"M{int(float(prio)):02d}"
            except (ValueError, TypeError):
                rid = f"M{idx:02d}"
        rules.append({"id": rid, "cells": cells,
                      "then": {"action": rec[ACTION_COL],
                               "notify": _split_list(rec.get(NOTIFY_COL, ""))}})

    return {"matrix": px["name"], "version": px["version"],
            "hit_policy": px["hit_policy"], "inputs": inputs_compiled,
            "outputs": [{"name": "action"}, {"name": "notify"}], "rules": rules}


def validate(px: dict) -> tuple[M.Matrix, list[str], list[str], list[tuple[str, str]]]:
    md = xlsx_to_matrix_dict(px)
    m = M.from_dict(md)
    errors, warnings, overlaps = M.lint(m)
    actions = set(px["actions"]) or set(DEFAULT_ACTIONS)
    for r in m.rules:
        if r.action not in actions:
            errors.append(f"{r.id}: 动作「{r.action}」不在动作枚举内")
        for t in r.notify:
            if px["notify_targets"] and t not in px["notify_targets"]:
                errors.append(f"{r.id}: 抄送对象「{t}」不在抄送枚举内")
    if px["mapping"]:
        for r in m.rules:
            if r.action not in px["mapping"]:
                errors.append(f"{r.id}: 动作「{r.action}」缺少审批人映射")
    return m, errors, warnings, overlaps


def export_to_xlsx(m: M.Matrix, path: str | Path, changelog: list[tuple] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_RULES
    ws["A1"] = f"matrix: {m.name} | hit_policy: {m.hit_policy}"
    headers = [PRIORITY_COL, ID_COL] + [(m.inputs[c].get("header") or c) for c in m.inputs] \
        + [ACTION_COL, NOTIFY_COL]
    ws.append(headers)
    for i, r in enumerate(m.rules, 1):
        ws.append([i, r.id,
                   *[(r.cells[c].raw if c in r.cells else M.ANY) for c in m.inputs],
                   r.action, ",".join(r.notify)])

    ws2 = wb.create_sheet(SHEET_ENUM)
    ws2.append([SEC_INPUT])
    ws2.append(["列名", "类型", "枚举", "表头"])
    for c, d in m.inputs.items():
        ws2.append([c, d["type"], ",".join(d["enum"]) if d.get("enum") else "",
                    d.get("header") or ""])
    ws2.append([])
    ws2.append([SEC_OUTPUT])
    ws2.append(["动作枚举", ", ".join(DEFAULT_ACTIONS)])
    ws2.append(["抄送对象", ", ".join(DEFAULT_NOTIFY)])
    ws2.append([])
    ws2.append([SEC_MAP])
    ws2.append(["AUTO_APPROVE", "-"])
    ws2.append(["MANAGER", "manager"])
    ws2.append(["MANAGER_BOSS", "manager,boss"])
    ws2.append(["MANAGER_BOSS_FINANCE_CONSIGN", "manager,boss,finance_consign"])
    ws2.append(["MANUAL_REVIEW", "boss"])

    ws3 = wb.create_sheet(SHEET_LOG)
    ws3.append(["version", "日期", "修改人", "变更摘要"])
    for row in (changelog or [(m.version, "", "export", "由 YAML 导出生成，可在 Sheet1 编辑后 import/commit")]):
        ws3.append(list(row))

    for w, col in ((10, "A"), (16, "B"), (18, "C"), (12, "D"), (12, "E"),
                   (30, "F"), (14, "G")):
        ws.column_dimensions[col].width = w
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 40
    wb.save(path)
    return Path(path)


def boundary_contexts(m: M.Matrix, cap: int = 5000) -> list[dict]:
    nums: set[float] = set()
    for r in m.rules:
        for cell in r.cells.values():
            if isinstance(cell, (M.CompareCell, M.RangeCell)):
                iv = cell.interval()
                for b, _ in (iv.lo, iv.hi):
                    if b not in (float("inf"), float("-inf")):
                        nums.update({b, b - 1, b + 1})
    cols: list[list[Any]] = []
    for c, d in m.inputs.items():
        if d["type"] == "number":
            cols.append(sorted(nums) + [None])
        elif d.get("enum"):
            cols.append(list(d["enum"]) + [None])
        else:
            cols.append(["__probe__", None])
    ctxs = [dict(zip(m.inputs.keys(), combo)) for combo in
            itertools.islice(itertools.product(*cols), cap)]
    return ctxs


def _evaluate_safe(m: M.Matrix, ctx: dict) -> tuple[str, str]:
    try:
        d = M.evaluate(m, ctx)
        return d.hit_rule_id or "-", d.action
    except M.NoMatchError:
        return "NO_MATCH", M.FALLBACK_ACTION
    except M.UniqueViolationError:
        return "UNIQUE_VIOLATION", M.FALLBACK_ACTION


def dry_run(new_m: M.Matrix, old_m: M.Matrix | None, contexts: list[dict]) -> str:
    new_dist: dict[str, int] = {}
    new_act: dict[str, int] = {}
    for ctx in contexts:
        rid, a = _evaluate_safe(new_m, ctx)
        new_dist[rid] = new_dist.get(rid, 0) + 1
        new_act[a] = new_act.get(a, 0) + 1
    lines = [f"干跑样本 {len(contexts)} 条", "", "命中分布（新矩阵 v%d）:" % new_m.version]
    for k in sorted(new_dist):
        lines.append(f"  {k:<14} {new_dist[k]:>6}")
    if old_m is not None:
        old_act: dict[str, int] = {}
        for ctx in contexts:
            _, a = _evaluate_safe(old_m, ctx)
            old_act[a] = old_act.get(a, 0) + 1
        lines += ["", "动作分布变化（old → new）："]
        for a in sorted(set(old_act) | set(new_act)):
            o, n = old_act.get(a, 0), new_act.get(a, 0)
            lines.append(f"  {a:<32} {o:>5} → {n:>5}{'' if o == n else '   ← 变化'}")
    return "\n".join(lines)


def commit(xlsx_path: str | Path, db_path: str | Path, operator: str = "unknown",
           audit_dir: str | Path = "data/audit") -> int:
    px = read_xlsx(xlsx_path)
    m, errors, _, _ = validate(px)
    if errors:
        raise M.MatrixLintError("校验失败: " + "; ".join(errors))
    name, version = m.name, m.version
    src = Path(xlsx_path)
    checksum = hashlib.sha256(src.read_bytes()).hexdigest()
    yaml_text = yaml.safe_dump(xlsx_to_matrix_dict(px), allow_unicode=True, sort_keys=False)

    conn = MD.init_db(db_path)
    prev = conn.execute(
        "SELECT version, compiled_yaml FROM matrix_versions"
        " WHERE matrix_name=? AND status='active' ORDER BY version DESC LIMIT 1",
        (name,)).fetchone()
    if prev and version <= prev[0]:
        raise M.MatrixLintError(
            f"版本必须递增：当前 active v{prev[0]}，导入 v{version}（请在 Sheet3_变更记录 升版本）")
    conn.execute("UPDATE matrix_versions SET status='retired'"
                 " WHERE matrix_name=? AND status='active'", (name,))
    conn.execute(
        "INSERT INTO matrix_versions(matrix_name, version, source_xlsx, checksum,"
        " compiled_yaml, hit_policy, status, created_by)"
        " VALUES(?,?,?,?,?,?,'active',?)",
        (name, version, str(src), checksum, yaml_text, m.hit_policy, operator))

    diff = ""
    if prev:
        diff = "".join(difflib.unified_diff(
            (prev[1] or "").splitlines(True), yaml_text.splitlines(True),
            fromfile=f"{name} v{prev[0]}", tofile=f"{name} v{version}"))
    adir = Path(audit_dir)
    adir.mkdir(parents=True, exist_ok=True)
    diff_file = adir / f"matrix_{name}_v{version}.diff"
    diff_file.write_text(diff or "(无上一版本，首次激活)", encoding="utf-8")
    MD.append_audit(conn, ticket_id=None, actor=operator, action="matrix.import",
                    payload_hash=hashlib.sha256(diff_file.read_bytes()).hexdigest(),
                    payload_file=str(diff_file))
    conn.commit()
    conn.close()
    return version
