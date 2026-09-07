"""月度统计报告（PRD F16）：数据聚合 → matplotlib 图表 → docx/pptx 双格式。

口径（诚实版）：单据状态为**当前快照**（tickets 无状态历史表），"本月新增=occurred_at
归属月的全部单据"；金额统计排除驳回/撤回/作废；图表中文字体自动适配 macOS/Windows。
CLI: boss-report --db data/secretary.db [--month 2026-09] [--format docx|pptx|both]
"""
from __future__ import annotations

import datetime as dt
import sqlite3
from pathlib import Path
from typing import Any, Mapping, Sequence

COUNTED = ("APPROVED", "PAID", "AUTO_APPROVED", "SUBMITTED", "ESCALATED")


def last_completed_month(today: dt.date | None = None) -> str:
    today = today or dt.date.today()
    first = today.replace(day=1)
    return (first - dt.timedelta(days=1)).strftime("%Y-%m")


def _setup_font() -> None:
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import font_manager, rcParams
    names = {"PingFang SC", "Hiragino Sans GB", "Microsoft YaHei", "SimHei",
             "Noto Sans CJK SC", "WenQuanYi Micro Hei"}
    for cand in names:
        if any(cand.lower() in (f.name or "").lower()
               for f in font_manager.fontManager.ttflist):
            rcParams["font.sans-serif"] = [cand]
            break
    rcParams["axes.unicode_minus"] = False


def _ticket_rows(conn) -> list[dict]:
    cols = ("ticket_id", "employee_id", "dept_id", "status", "amount", "currency",
            "expense_type", "reason", "occurred_at", "created_at", "submitted_at",
            "allowance_id")
    rows = conn.execute(f"SELECT {','.join(cols)} FROM tickets").fetchall()
    return [dict(zip(cols, r)) for r in rows]


def collect(conn, month: str) -> dict:
    now = dt.datetime.now()
    tickets = _ticket_rows(conn)
    month_t = [t for t in tickets if (t.get("occurred_at") or "")[:7] == month]

    def amt(ts, statuses=COUNTED):
        return round(sum(float(t.get("amount") or 0) for t in ts
                         if t.get("status") in statuses), 2)

    trend = []
    base = dt.date.fromisoformat(month + "-01")
    for i in range(5, -1, -1):
        m = (base - dt.timedelta(days=30 * i)).strftime("%Y-%m") if i else month
        if i:
            m = _shift_month(base, -i)
        sub = [t for t in tickets if (t.get("occurred_at") or "")[:7] == m]
        trend.append({"month": m, "amount": amt(sub), "count": len(sub)})
    type_sum: dict[str, float] = {}
    for t in month_t:
        if t.get("status") in COUNTED:
            k = t.get("expense_type") or "其他"
            type_sum[k] = type_sum.get(k, 0) + float(t.get("amount") or 0)
    dept_sum: dict[str, float] = {}
    for t in month_t:
        if t.get("status") in COUNTED:
            k = t.get("dept_id") or "*"
            dept_sum[k] = dept_sum.get(k, 0) + float(t.get("amount") or 0)
    status_counts: dict[str, int] = {}
    for t in month_t:
        status_counts[t.get("status")] = status_counts.get(t.get("status"), 0) + 1
    anomalies = [dict(zip(("type", "subject", "period", "severity", "evidence"), r))
                 for r in conn.execute(
        "SELECT type, subject, period, severity, evidence FROM anomalies"
        " WHERE period=? ORDER BY severity DESC", (month,)).fetchall()]
    allowances = [dict(zip(("allowance_id", "employee_id", "category", "total_amount",
                            "used_amount", "status"), r)) for r in conn.execute(
        "SELECT allowance_id, employee_id, category, total_amount, used_amount, status"
        " FROM allowances WHERE status IN ('active','EXHAUSTED')").fetchall()]
    allowance_used_n = sum(1 for t in month_t if t.get("allowance_id"))
    budgets = [dict(zip(("dept", "amount"), r)) for r in conn.execute(
        "SELECT dept_id, amount FROM budgets WHERE month=?", (month,)).fetchall()]
    return {"month": month, "generated_at": now.isoformat(timespec="minutes"),
            "tickets": tickets, "month_count": len(month_t),
            "amount": amt(month_t), "paid_amount": round(
                sum(float(t.get("amount") or 0) for t in month_t
                    if t.get("status") == "PAID"), 2),
            "rejected_n": sum(1 for t in month_t if t.get("status") == "REJECTED"),
            "pending_n": sum(1 for t in month_t
                             if t.get("status") in ("SUBMITTED", "ESCALATED")),
            "direct_approved_n": sum(1 for t in month_t
                                     if t.get("status") == "AUTO_APPROVED"),
            "allowance_used_n": allowance_used_n,
            "trend": trend, "type_sum": type_sum, "dept_sum": dept_sum,
            "status_counts": status_counts, "anomalies": anomalies,
            "allowances": allowances, "budgets": budgets}


def _shift_month(base: dt.date, delta: int) -> str:
    y, m = base.year, base.month + delta
    y += (m - 1) // 12
    m = (m - 1) % 12 + 1
    return f"{y}-{m:02d}"


def build_charts(data: Mapping, outdir: Path) -> list[Path]:
    _setup_font()
    import matplotlib.pyplot as plt
    outdir.mkdir(parents=True, exist_ok=True)
    paths = []

    # 图1 近6个月趋势（柱=金额, 线=单量）
    months = [t["month"] for t in data["trend"]]
    amounts = [t["amount"] for t in data["trend"]]
    counts = [t["count"] for t in data["trend"]]
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.bar(months, amounts, color="#4C78A8", label="金额(元)")
    ax2 = ax.twinx()
    ax2.plot(months, counts, color="#E45756", marker="o", label="单量")
    ax.set_ylabel("金额(元)")
    ax2.set_ylabel("单量")
    ax.set_title(f"近6个月报销趋势（{data['month']}）")
    for a in (ax, ax2):
        a.tick_params(axis="x", labelrotation=30)
    fig.legend(loc="upper left", bbox_to_anchor=(0.08, 0.95))
    fig.tight_layout()
    p = outdir / f"trend_{data['month']}.png"
    fig.savefig(p, dpi=150)
    plt.close(fig)
    paths.append(p)

    # 图2 费用类型分布
    if data["type_sum"]:
        fig, ax = plt.subplots(figsize=(6, 4.5))
        ax.pie(list(data["type_sum"].values()),
               labels=list(data["type_sum"].keys()), autopct="%1.0f%%",
               startangle=90)
        ax.set_title(f"费用类型分布（{data['month']}）")
        fig.tight_layout()
        p = outdir / f"type_{data['month']}.png"
        fig.savefig(p, dpi=150)
        plt.close(fig)
        paths.append(p)

    # 图3 部门费用 vs 预算
    if data["dept_sum"]:
        fig, ax = plt.subplots(figsize=(7, 4))
        depts = list(data["dept_sum"].keys())
        used = [data["dept_sum"][d] for d in depts]
        budget_map = {b["dept"]: float(b["amount"]) for b in data["budgets"]}
        budgets = [budget_map.get(d, 0) for d in depts]
        x = range(len(depts))
        ax.bar([i - 0.2 for i in x], used, width=0.4, label="已用", color="#4C78A8")
        ax.bar([i + 0.2 for i in x], budgets, width=0.4, label="预算",
               color="#BAB0AC")
        ax.set_xticks(list(x))
        ax.set_xticklabels(depts)
        ax.set_title(f"部门费用 vs 预算（{data['month']}）")
        ax.legend()
        fig.tight_layout()
        p = outdir / f"dept_{data['month']}.png"
        fig.savefig(p, dpi=150)
        plt.close(fig)
        paths.append(p)

    # 图4 状态分布
    if data["status_counts"]:
        fig, ax = plt.subplots(figsize=(7, 3.5))
        ks = list(data["status_counts"].keys())
        ax.bar(ks, [data["status_counts"][k] for k in ks], color="#54A24B")
        ax.set_title(f"本月新增单据状态分布（当前快照）")
        fig.tight_layout()
        p = outdir / f"status_{data['month']}.png"
        fig.savefig(p, dpi=150)
        plt.close(fig)
        paths.append(p)
    return paths


def render_markdown(data: Mapping, images: Sequence[Path]) -> str:
    lines = [f"# 秘书月报 · {data['month']}",
             f"生成时间：{data['generated_at']}", "",
             f"本月单据 **{data['month_count']}** 张，金额 **{data['amount']}** 元；"
             f"已打款 {data['paid_amount']} 元；驳回 {data['rejected_n']} 单；"
             f"在途 {data['pending_n']} 单；AI 直批 {data['direct_approved_n']} 单；"
             f"额度核销 {data['allowance_used_n']} 单。", ""]
    for img in images:
        lines.append(f"![{img.stem}]({img})")
        lines.append("")
    if data["anomalies"]:
        lines.append("## 异常事件")
        lines.append("")
        lines.append("| 级别 | 类型 | 对象 | 说明 |")
        lines.append("|------|------|------|------|")
        for a in data["anomalies"]:
            lines.append(f"| {a['severity']} | {a['type']} | {a['subject']} | "
                         f"{str(a['evidence'])[:60]} |")
        lines.append("")
    if data["allowances"]:
        lines.append("## 额度使用")
        lines.append("")
        for a in data["allowances"]:
            used = float(a.get("used_amount") or 0)
            lines.append(f"- {a['allowance_id']} {a['employee_id']} "
                         f"{a['category']}：已用 {used:.0f}/{a['total_amount']:.0f} 元"
                         f"（{a['status']}）")
    lines += ["", "> 口径说明：单据状态为当前快照；金额统计排除驳回/撤回/作废。",
              "> 本报告仅供内部管理参考，不构成审计意见。"]
    return "\n".join(lines)


def export_docx(data: Mapping, images: Sequence[Path], out: Path) -> Path:
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading(f"秘书月报 · {data['month']}", 0)
    doc.add_paragraph(f"生成时间：{data['generated_at']}")
    doc.add_paragraph(
        f"本月单据 {data['month_count']} 张，金额 {data['amount']} 元；"
        f"已打款 {data['paid_amount']} 元；驳回 {data['rejected_n']} 单；"
        f"在途 {data['pending_n']} 单；AI 直批 {data['direct_approved_n']} 单；"
        f"额度核销 {data['allowance_used_n']} 单。")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "类型", "金额(元)", "占比"
    total = data["amount"] or 1
    for k, v in data["type_sum"].items():
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = k, f"{v:.2f}", \
            f"{v / total * 100:.0f}%"
    for img in images:
        doc.add_picture(str(img), width=Inches(6))
    if data["anomalies"]:
        doc.add_heading("异常事件", level=1)
        for a in data["anomalies"]:
            doc.add_paragraph(f"[{a['severity']}] {a['type']} {a['subject']}："
                              f"{str(a['evidence'])[:80]}")
    doc.add_paragraph("口径说明：单据状态为当前快照；金额统计排除驳回/撤回/作废。"
                      "本报告仅供内部管理参考，不构成审计意见。")
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def export_pptx(data: Mapping, images: Sequence[Path], out: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = f"秘书月报 · {data['month']}"
    s.placeholders[1].text = (f"生成 {data['generated_at']} | "
                              f"单据 {data['month_count']} 张 / {data['amount']} 元")
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "核心指标"
    tf = s2.placeholders[1].text_frame
    for line in (
        f"本月金额：{data['amount']} 元（已打款 {data['paid_amount']}）",
        f"AI 直批：{data['direct_approved_n']} 单 | 额度核销：{data['allowance_used_n']} 单",
        f"驳回 {data['rejected_n']} 单 | 在途 {data['pending_n']} 单",
        f"费用构成：{', '.join(f'{k} {v:.0f}元' for k, v in data['type_sum'].items()) or '-'}",
    ):
        tf.add_paragraph().text = line
    for img in images:
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.add_picture(str(img), Inches(1.5), Inches(1.2), width=Inches(10.3))
    if data["anomalies"]:
        s3 = prs.slides.add_slide(prs.slide_layouts[1])
        s3.shapes.title.text = "异常事件"
        tf3 = s3.placeholders[1].text_frame
        tf3.text = f"共 {len(data['anomalies'])} 条（详见 docx）"
        for a in data["anomalies"][:5]:
            tf3.add_paragraph().text = f"[{a['severity']}] {a['subject']}"
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def generate(conn, month: str | None = None, outdir: str | Path = "data/reports",
             fmt: str = "both") -> dict:
    month = month or last_completed_month()
    outdir = Path(outdir)
    data = collect(conn, month)
    images = build_charts(data, outdir)
    result = {"month": month, "images": [str(p) for p in images]}
    if fmt in ("docx", "both"):
        result["docx"] = str(export_docx(data, images,
                                         outdir / f"秘书月报_{month}.docx"))
    if fmt in ("pptx", "both"):
        result["pptx"] = str(export_pptx(data, images,
                                         outdir / f"秘书月报_{month}.pptx"))
    result["markdown"] = render_markdown(data, images)
    return result


def main(argv: list[str] | None = None) -> int:
    import argparse
    p = argparse.ArgumentParser(prog="boss_secretary.report.monthly")
    p.add_argument("--db", default="data/secretary.db")
    p.add_argument("--month", default=None, help="默认上个完整月")
    p.add_argument("--format", default="both", choices=["docx", "pptx", "both"])
    p.add_argument("--out", default="data/reports")
    args = p.parse_args(argv)
    import sqlite3
    conn = sqlite3.connect(args.db, check_same_thread=False)
    result = generate(conn, month=args.month, outdir=args.out, fmt=args.format)
    print("月报生成完成：")
    for k, v in result.items():
        if k != "markdown":
            print(f"  {k}: {v}")
    return 0


if __name__ == "__main__":
    import sys
    sys.exit(main())
